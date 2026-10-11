"""Generate template_cbt.docx with Jinja2 tags for docxtpl.

Run once:  python make_template.py
Produces:  template_cbt.docx — tabel Thought Record 5 kolom
           ({{ situation }}, {{ thought }}, {{ emotion }}, {{ behavior }},
            {{ alternate_thought }})
"""

from pathlib import Path

from docx import Document
from docx.shared import Pt

OUT = Path(__file__).parent / "template_cbt.docx"

HEADERS = ["Situasi", "Pikiran", "Emosi", "Perilaku", "Pikiran Alternatif"]
TAGS = ["situation", "thought", "emotion", "behavior", "alternate_thought"]


def main() -> None:
    doc = Document()

    h = doc.add_heading("Thought Record — Matriks CBT", level=0)
    h.alignment = 1

    doc.add_paragraph("Dokumen dibuat otomatis oleh HearMeOut (MCP: HearMeOut_Tools).")

    table = doc.add_table(rows=2, cols=len(HEADERS))
    table.style = "Table Grid"

    header_cells = table.rows[0].cells
    for cell, title in zip(header_cells, HEADERS):
        run = cell.paragraphs[0].add_run(title)
        run.bold = True
        run.font.size = Pt(11)

    body_cells = table.rows[1].cells
    for cell, tag in zip(body_cells, TAGS):
        run = cell.paragraphs[0].add_run("{{ %s }}" % tag)
        run.font.size = Pt(11)

    doc.add_paragraph("")
    note = doc.add_paragraph(
        "Catatan: dokumen ini bersifat objektif-klinis, berisi ekstraksi data saja "
        "tanpa saran medis. Teks bersumber dari jurnal pasien (tanpa data pribadi)."
    )
    note.runs[0].italic = True

    doc.save(OUT)
    print(f"Template written: {OUT}")


if __name__ == "__main__":
    main()
